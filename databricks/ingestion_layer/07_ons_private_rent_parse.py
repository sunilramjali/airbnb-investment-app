# ONS Price Index of Private Rents (PIPR) — Bronze xlsx parser.
# Databricks port of the Python stored procedure in etl/ingestion_layer/07_ons_private_rent_ddl.sql.
# ============================================================
# ONS PRIVATE RENT (PIPR)  —  PARSE (imported by 08_ons_private_rent_load.py).
# ------------------------------------------------------------
# SOURCE / DESIGN:
#   A monthly Lambda (ons_pipr) downloads the ONS workbook UNCHANGED and lands the
#   raw spreadsheet at:
#     s3://airbnb-investment-app-988261629236-eu-west-2-an/raw/ons/private-rents/<file>.xlsx
#   This module parses the NEWEST .xlsx with openpyxl and appends one faithful row
#   per spreadsheet row (all cells as an ARRAY of TEXT).
#
# ============================================================
# WHY THIS IS A .py AND NOT A .sql
# ------------------------------------------------------------
# The Snowflake original was a CREATE PROCEDURE ... LANGUAGE PYTHON living inside a
# .sql file and invoked with CALL. Databricks has no SQL-invoked Python stored
# procedure of that shape, so the parse becomes an ordinary Python function called
# from a notebook/job task. Keeping the .sql extension would misrepresent the file.
#
# This is the one genuine INVOCATION-MODEL change in the Bronze port: everything
# else translated construct-for-construct. The parse LOGIC below is otherwise a
# faithful copy of the proc — same sheet picking, same skip-rows behaviour, same
# original-row-number preservation.
#
# ============================================================
# SNOWFLAKE -> DATABRICKS TRANSLATION NOTES
# ------------------------------------------------------------
#   CREATE STAGE ... DIRECTORY = (ENABLE = TRUE)
#       -> no counterpart. The external location `airbnb_raw_ons` plus a LIST
#          statement replaces both the stage and its directory table. There is no
#          ALTER STAGE ... REFRESH to run: LIST always reads live from S3, which
#          removes a whole class of stale-directory bug the Snowflake version had
#          to defend against.
#
#   SELECT RELATIVE_PATH FROM DIRECTORY(@STAGE) ORDER BY LAST_MODIFIED DESC
#       -> LIST '<s3 path>'  then sort on the modification_time column.
#
#   session.file.get('@STAGE/<f>', '/tmp')  +  load_workbook('/tmp/<f>')
#       -> spark.read.format('binaryFile') + io.BytesIO. Serverless compute has no
#          durable local filesystem to stage a download into, and openpyxl accepts
#          a file-like object, so the bytes go straight from S3 into the parser
#          without ever touching disk.
#
#   session.create_dataframe(records, schema).write.save_as_table(...)
#       -> spark.createDataFrame(records, schema).write.saveAsTable(...)
#
#   ⚠️ _FILE_ROW_NUMBER HERE IS NOT THE SYNTHETIC ONE USED ELSEWHERE IN BRONZE.
#      In 02_bronze_load.py and 04_land_registry_load.sql, _FILE_ROW_NUMBER is a
#      synthetic row_number() standing in for Snowflake's METADATA$FILE_ROW_NUMBER.
#      Here it is the REAL 1-based sheet row index from enumerate(), exactly as in
#      the Snowflake proc — and it is load-bearing: Silver filters
#      `WHERE _FILE_ROW_NUMBER >= 4` to skip the ONS title+header rows
#      (etl/cleaning_layer/13_silver_ons_private_rent.sql:56). It must keep
#      counting original sheet rows, including the ones skipped by SKIP_ROWS.
# ============================================================

import io

from pyspark.sql.types import (
    ArrayType, LongType, StringType, StructField, StructType,
)

# Sheets that are never the data tab.
SKIP_SHEETS = {"contents", "notes", "cover", "metadata"}


def _pick_sheet(wb, requested):
    """Choose the workbook tab to parse (tolerant of case / trailing spaces)."""
    if requested:
        want = requested.strip().lower()
        for name in wb.sheetnames:
            if name.strip().lower() == want:
                return wb[name]
        raise ValueError(f"Sheet '{requested}' not found. Tabs: {wb.sheetnames}")
    for name in wb.sheetnames:
        if name.strip().lower() not in SKIP_SHEETS:
            return wb[name]
    return wb[wb.sheetnames[0]]


def newest_xlsx(session, prefix: str) -> str:
    """Return the full S3 path of the most recently modified .xlsx under `prefix`.

    Replaces the Snowflake proc's ALTER STAGE REFRESH + DIRECTORY() query. LIST reads
    live from S3, so there is no directory table to keep in sync.
    """
    rows = session.sql(f"LIST '{prefix}'").collect()
    xlsx = [r for r in rows if str(r["path"]).lower().endswith(".xlsx")]
    if not xlsx:
        raise FileNotFoundError(f"No .xlsx files found under {prefix}")
    newest = max(xlsx, key=lambda r: r["modification_time"])
    return str(newest["path"])


def _read_bytes(session, path: str) -> bytes:
    """Pull the workbook's bytes through Spark rather than a local download.

    Serverless compute has no durable local disk for session.file.get()'s /tmp
    staging, and binaryFile is the supported way to hand raw bytes to a Python
    library.
    """
    row = session.read.format("binaryFile").load(path).select("content").collect()[0]
    return row["content"]


def _audit(session, audit_table, file_name, parsed, loaded, status, err=None):
    """Write one LOAD_AUDIT row. A Python load has no COPY history to read back,
    exactly as in the Snowflake proc, so the loader records its own outcome."""
    session.sql(
        f"""
        INSERT INTO {audit_table}
            (TABLE_NAME, FILE_NAME, STATUS, ROWS_PARSED, ROWS_LOADED,
             ERRORS_SEEN, FIRST_ERROR, FIRST_ERROR_LINE)
        VALUES ('RAW_ONS_PRIVATE_RENT', :fname, :status, :parsed, :loaded,
                :errors, :err, NULL)
        """,
        args={
            "fname": file_name,
            "status": status,
            "parsed": parsed,
            "loaded": loaded,
            "errors": 0 if err is None else 1,
            "err": err,
        },
    )


SCHEMA = StructType([
    StructField("SHEET", StringType()),
    StructField("CELLS", ArrayType(StringType())),
    StructField("_FILENAME", StringType()),
    StructField("_FILE_ROW_NUMBER", LongType()),
])


def load_ons_private_rent(session, prefix, target_table, audit_table,
                          sheet_name="Table 1", skip_rows=2) -> str:
    """Parse the newest ONS PIPR workbook and append faithful rows to `target_table`.

    SHEET_NAME : workbook tab to read (matched case-insensitively / trimmed);
                 None -> auto-pick the first non-Contents/Notes sheet.
    SKIP_ROWS  : number of leading rows to drop (ONS puts 2 title rows above the
                 header). _FILE_ROW_NUMBER keeps the ORIGINAL sheet row number, so
                 the header row stays identifiable downstream.
    """
    from openpyxl import load_workbook   # imported here so the module loads without it

    skip_rows = int(skip_rows or 0)

    path = newest_xlsx(session, prefix)
    wb = load_workbook(io.BytesIO(_read_bytes(session, path)),
                       read_only=True, data_only=True)
    ws = _pick_sheet(wb, sheet_name)

    records = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= skip_rows:                      # drop leading title rows
            continue
        if all(c is None for c in row):
            continue
        cells = [None if c is None else str(c) for c in row]
        records.append([ws.title, cells, path, i])   # keep ORIGINAL row number

    if not records:
        _audit(session, audit_table, path, 0, 0, "LOADED")
        return f"File {path}: no data rows after skipping {skip_rows}."

    (session.createDataFrame(records, SCHEMA)
            .write.mode("append")
            .saveAsTable(target_table))

    _audit(session, audit_table, path, len(records), len(records), "LOADED")
    return (
        f"Loaded {len(records)} rows from {path} "
        f"(sheet {ws.title}, skipped first {skip_rows} rows)."
    )
